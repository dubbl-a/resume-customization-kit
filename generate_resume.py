#!/usr/bin/env python3
"""
generate_resume.py — Generate a formatted .docx resume from a JSON content file.

Uses the MASTER resume as a style template so output exactly matches formatting.

Usage:
    python3 generate_resume.py content.json "Your Name Resume MM.DD.YYYY Company Role.docx"
    python3 generate_resume.py content.json "Output.docx" --pages 1    # one-page resume
    python3 generate_resume.py content.json "Output.docx" --no-pdf     # .docx only

After exporting the PDF (Word first, LibreOffice as fallback) the script counts
its pages and exits 1 unless the count matches --pages (default 2), or if no PDF
could be produced.

Master resume lookup, first match wins:
    1. MASTER_RESUME environment variable (full path to a .docx)
    2. master_resume.docx next to this script
    3. Exactly one *MASTER*.docx next to this script (Word lock files ~$... ignored)

JSON format (see resume_instructions.md for field descriptions):
{
  "headline": "Marketing Leader | Integrated GTM, Brand & Product Storytelling",
  "summary": "Two-sentence summary with **bold anchors** inline.",
  "experience": [
    {
      "company": "ADOBE",
      "title": "Sr GTM Strategy Manager, Express",
      "location": "Los Angeles, CA",
      "dates": "Aug 2025-Present",
      "bullets": [
        "**Topic Header**: Body text with **bold metric** inline."
      ]
    }
  ],
  "skills": [
    { "header": "Integrated Marketing", "items": "GTM Strategy, Launch Planning, ..." }
  ],
  "education": "Bachelor of Arts, Economics (Cum Laude honors) | Tufts University, Medford, MA"
}
"""

import sys
import json
import os
import re
import glob
import argparse
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Constants ────────────────────────────────────────────────────────────────

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = SCRIPT_DIR

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# Fallback right-aligned tab stop: Letter page (12240) minus margins (1152 each) = 9936 DXA.
# Normally computed from the master's page setup (see right_tab_pos).
DEFAULT_RIGHT_TAB_POS = "9936"
RIGHT_TAB_POS = DEFAULT_RIGHT_TAB_POS

# Paragraph styles the master must define
REQUIRED_STYLES = ("Body A", "Section Title", "List Paragraph")


# ── Master resume lookup ──────────────────────────────────────────────────────

def _find_master():
    """
    Locate the master resume .docx used as the style template:
      1. MASTER_RESUME environment variable
      2. master_resume.docx next to this script
      3. Exactly one *MASTER*.docx next to this script (case-insensitive,
         Word lock files starting with ~$ ignored)
    Exits with a clear message if nothing is found or the choice is ambiguous.
    """
    env = os.environ.get("MASTER_RESUME")
    if env:
        if os.path.isfile(env):
            return env
        print(f"ERROR: MASTER_RESUME is set but the file does not exist: {env}")
        sys.exit(1)

    default = os.path.join(SCRIPT_DIR, "master_resume.docx")
    if os.path.isfile(default):
        return default

    candidates = sorted(
        f for f in glob.glob(os.path.join(SCRIPT_DIR, "*.docx"))
        if "MASTER" in os.path.basename(f).upper()
        and not os.path.basename(f).startswith("~$")
    )
    if len(candidates) == 1:
        return candidates[0]

    if not candidates:
        print(f"ERROR: No master resume .docx found in {SCRIPT_DIR}")
    else:
        print(f"ERROR: Found {len(candidates)} *MASTER*.docx files; expected exactly one:")
        for f in candidates:
            print(f"  - {os.path.basename(f)}")
    print("  Option 1: Name your master resume master_resume.docx (checked first)")
    print("  Option 2: Keep exactly one .docx with 'MASTER' in its filename")
    print("  Option 3: Set MASTER_RESUME to the full path of your master .docx")
    sys.exit(1)


# ── Master preflight ──────────────────────────────────────────────────────────

def preflight_master(doc, master_path):
    """
    Check that the master has the structure this script writes into:
      - Paragraph styles "Body A", "Section Title" and "List Paragraph"
      - A header table whose first cell has at least two paragraphs, where
        paragraph 1 holds the headline, a <w:br/> line break, then contact info
    Lists every problem and exits 1 instead of failing mid-build.
    """
    problems = []

    style_names = {s.name for s in doc.styles}
    for name in REQUIRED_STYLES:
        if name not in style_names:
            problems.append(f'missing paragraph style "{name}"')

    if not doc.tables:
        problems.append("no header table (the name/headline/contact block must be "
                        "the first table in the document)")
    else:
        cell = doc.tables[0].rows[0].cells[0]
        paras = cell.paragraphs
        if len(paras) < 2:
            problems.append(f"header table's first cell has {len(paras)} paragraph(s); "
                            "expected name, then headline + line break + contact")
        elif not any(child.tag == qn('w:r') and child.find(qn('w:br')) is not None
                     for child in paras[1]._p):
            problems.append("header table's second paragraph has no line break (<w:br/>) "
                            "separating the headline from the contact line")

    if problems:
        print("ERROR: Master resume is missing structure this script needs:")
        print(f"  {master_path}")
        for p in problems:
            print(f"  - {p}")
        print("  See README.md (Add your master resume) for the required structure.")
        sys.exit(1)


def right_tab_pos(doc):
    """Right margin position in twips: page width minus left and right margins."""
    section = doc.sections[-1]
    if None in (section.page_width, section.left_margin, section.right_margin):
        return DEFAULT_RIGHT_TAB_POS
    usable = section.page_width - section.left_margin - section.right_margin
    return str(int(round(usable / 635)))  # 635 EMU per twip


# ── Low-level XML helpers ─────────────────────────────────────────────────────

def get_or_create_pPr(para):
    pPr = para._element.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr')
        para._element.insert(0, pPr)
    return pPr


def get_or_create_rPr(run):
    rPr = run._element.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        run._element.insert(0, rPr)
    return rPr


def set_run_sz(run, sz_val):
    """Set explicit font size (in half-points) on a run."""
    rPr = get_or_create_rPr(run)
    for tag in ('w:sz', 'w:szCs'):
        el = OxmlElement(tag)
        el.set(qn('w:val'), str(sz_val))
        rPr.append(el)


def add_tab_run(para, sz_val=24):
    """Append a <w:r><w:rPr><w:sz/></w:rPr><w:tab/></w:r> to a paragraph."""
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(sz_val))
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), str(sz_val))
    rPr.append(sz)
    rPr.append(szCs)
    r.append(rPr)
    tab = OxmlElement('w:tab')
    r.append(tab)
    para._p.append(r)


def add_line_break(para):
    """Append a <w:r><w:br/></w:r> to a paragraph (soft line break within bullet)."""
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    r.append(br)
    para._p.append(r)


# ── Bold-marker parser ────────────────────────────────────────────────────────

def parse_bold(text):
    """
    Parse **bold** markers. Returns list of (segment_text, is_bold) tuples.
    Unclosed markers treated as plain text.
    """
    parts = []
    current = ""
    in_bold = False
    i = 0
    while i < len(text):
        if text[i:i+2] == "**":
            if current:
                parts.append((current, in_bold))
                current = ""
            in_bold = not in_bold
            i += 2
        else:
            current += text[i]
            i += 1
    if current:
        parts.append((current, in_bold))
    if in_bold:
        flat = "".join(seg for seg, _ in parts)
        return [(flat, False)]
    return parts


def add_runs(para, text, base_bold=False, sz=None):
    """
    Add runs with **bold** markers parsed.
    IMPORTANT: never set run.bold = False — that writes <w:b w:val="0"/> which
    overrides the paragraph style's bold, breaking Section Title etc.
    """
    for seg, is_bold in parse_bold(text):
        run = para.add_run(seg)
        if base_bold or is_bold:
            run.bold = True
        if sz is not None:
            set_run_sz(run, sz)


# ── Paragraph builder ─────────────────────────────────────────────────────────

def add_para(doc_element, text, style_name, base_bold=False):
    """
    Append a paragraph with style-specific handling:

    Section Title  — explicit sz=24, bold on all runs (style doesn't define size)
    List Paragraph — injects <w:numPr> for bullet marker (numId=2)
                     also splits on ' – ' or ' - ' (standalone dash/en-dash)
                     and inserts soft line breaks instead of connector dashes
    All others     — plain add_runs
    """
    para = doc_element.add_paragraph(style=style_name)

    if style_name == "Section Title":
        add_runs(para, text, base_bold=True, sz=24)

    elif style_name == "List Paragraph":
        pPr = get_or_create_pPr(para)
        numPr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl')
        ilvl.set(qn('w:val'), '0')
        numId = OxmlElement('w:numId')
        numId.set(qn('w:val'), '2')
        numPr.append(ilvl)
        numPr.append(numId)
        pPr.insert(0, numPr)
        # Explicitly set justification so both Word and LibreOffice honour it.
        # Without this, LibreOffice ignores the style-level JUSTIFY and renders
        # bullets as ragged-right even though Word shows them correctly.
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), 'both')
        pPr.append(jc)
        # Split on standalone en dash or hyphen connectors → soft line break
        segments = re.split(r' [–\-] ', text)
        for i, segment in enumerate(segments):
            add_runs(para, segment, base_bold=base_bold)
            if i < len(segments) - 1:
                add_line_break(para)

    else:
        add_runs(para, text, base_bold=base_bold)

    return para


# ── Header — minimal in-place update ─────────────────────────────────────────

def update_headline_in_place(doc, headline):
    """
    Update only the headline text in the existing MASTER header table.
    Header cell paragraphs:
      0: Name
      1: Headline + <br/> + Contact (Body A, centered)
      2: Links (Body A, centered)
    Only touches paragraph 1: removes runs before <w:br/> run, inserts new bold headline run.
    Everything else (name, contact, links, hyperlinks, sizes) is left untouched.
    """
    cell = doc.tables[0].rows[0].cells[0]
    hl_para = cell.paragraphs[1]
    p = hl_para._p

    # Find the <w:r> that contains <w:br/> — contact line starts here
    br_parent_r = None
    for child in list(p):
        if child.tag == qn('w:r'):
            if child.find(qn('w:br')) is not None:
                br_parent_r = child
                break

    # Remove all <w:r> elements before br_parent_r
    for child in list(p):
        if child is br_parent_r:
            break
        if child.tag == qn('w:r'):
            p.remove(child)

    # Insert new bold headline run immediately before br_parent_r
    new_r = OxmlElement('w:r')
    new_rPr = OxmlElement('w:rPr')
    b = OxmlElement('w:b')
    new_rPr.append(b)
    new_r.append(new_rPr)
    new_t = OxmlElement('w:t')
    new_t.text = headline
    new_r.append(new_t)

    if br_parent_r is not None:
        br_parent_r.addprevious(new_r)
    else:
        p.append(new_r)


# ── Section builders ──────────────────────────────────────────────────────────

def build_summary(doc, summary_text):
    add_para(doc, "Summary", "Section Title")
    add_para(doc, summary_text, "Body A")


def add_role_spacer(doc):
    """
    Add a blank 6pt Normal paragraph before a role line.
    Matches the hand-built pattern: blank para with sz=12 in pPr rPr.
    """
    para = doc.add_paragraph()  # Normal style
    pPr = get_or_create_pPr(para)
    rPr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '12')
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), '12')
    rPr.append(sz)
    rPr.append(szCs)
    pPr.append(rPr)
    return para


def add_role_line(doc, company, title, location, dates):
    """
    Build a company/role line matching hand-built resume formatting:
      COMPANY (bold sz=24), , (normal sz=24), Title (italic sz=24),
      " | " + location (normal sz=24), one right-aligned tab, then dates.

    Right-aligned tab stop at RIGHT_TAB_POS forces dates flush to right margin
    regardless of text length.
    keepNext prevents orphaned role headers at page breaks.
    """
    para = doc.add_paragraph(style=doc.styles["Body A"])

    # Right-aligned tab stop + keepNext in paragraph properties
    pPr = get_or_create_pPr(para)

    keep_next = OxmlElement('w:keepNext')
    pPr.append(keep_next)

    tabs_el = OxmlElement('w:tabs')
    tab_el = OxmlElement('w:tab')
    tab_el.set(qn('w:val'), 'right')
    tab_el.set(qn('w:pos'), RIGHT_TAB_POS)
    tabs_el.append(tab_el)
    pPr.append(tabs_el)

    def sized_run(text, bold=False, italic=False):
        run = para.add_run(text)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        set_run_sz(run, 24)
        return run

    if company:
        sized_run(company, bold=True)
        sized_run(", ")
    sized_run(title, italic=True)
    sized_run(f" | {location}")
    # Single right-aligned tab jumps to the right margin
    add_tab_run(para, sz_val=24)
    # Dates — normal weight, sz=22 (inherits from BodyA style)
    para.add_run(dates)
    return para


def build_experience(doc, experience):
    add_para(doc, "professional experience", "Section Title")

    for i, role in enumerate(experience):
        add_role_spacer(doc)
        add_role_line(
            doc,
            role.get("company", ""),
            role.get("title", ""),
            role.get("location", ""),
            role.get("dates", ""),
        )
        for bullet in role.get("bullets", []):
            add_para(doc, bullet, "List Paragraph")


def build_skills(doc, skills):
    add_para(doc, "Skills", "Section Title")
    for skill in skills:
        add_para(doc, f"**{skill.get('header', '')}**: {skill.get('items', '')}", "Body A")


def build_education(doc, education):
    add_para(doc, "Education", "Section Title")
    add_para(doc, education, "Body A")


# ── Document settings ─────────────────────────────────────────────────────────

def disable_auto_hyphenation(doc):
    """
    Remove <w:autoHyphenation/> from document settings.
    The MASTER has this enabled, which causes Word to silently break words
    mid-syllable (Prod-uct, Strat-egy) on justified lines. Removing it
    prevents all automatic hyphenation.
    """
    settings = doc.settings.element
    for el in settings.findall(f"{{{W}}}autoHyphenation"):
        settings.remove(el)


# ── Body cleaner ──────────────────────────────────────────────────────────────

def clear_document_body(doc):
    """
    Remove all body paragraphs while preserving:
      - All <w:tbl> elements (header table must survive)
      - <w:sectPr> (page margins / section properties)
    """
    body = doc.element.body
    sect_pr = body.find(f"{{{W}}}sectPr")
    tables = set(body.findall(f"{{{W}}}tbl"))
    preserve = tables | {sect_pr}
    for child in list(body):
        if child not in preserve:
            body.remove(child)


# ── Font normalization ────────────────────────────────────────────────────────

def _replace_font(doc, old_font, new_font):
    """Replace every occurrence of old_font with new_font across the document XML.

    Used to swap Word-bundled Garamond for EB Garamond (a system-installed
    Garamond-family font accessible to both Word and LibreOffice), ensuring the
    DOCX and LibreOffice-generated PDF look identical.
    """
    W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    rFonts_tag = f"{{{W_NS}}}rFonts"
    font_attrs = [
        f"{{{W_NS}}}ascii",
        f"{{{W_NS}}}hAnsi",
        f"{{{W_NS}}}eastAsia",
        f"{{{W_NS}}}cs",
    ]
    for tree in (doc.element, doc.styles.element):
        for el in tree.iter(rFonts_tag):
            for attr in font_attrs:
                if el.get(attr) == old_font:
                    el.set(attr, new_font)


# ── PDF export ───────────────────────────────────────────────────────────────

def _applescript_str(path):
    """Escape a path for use inside an AppleScript string literal."""
    return path.replace('\\', '\\\\').replace('"', '\\"')


def _export_word_mac(docx_abs, pdf_abs):
    """Word via direct osascript: no file picker, no docx2pdf dialogs."""
    import subprocess, shutil

    if sys.platform != "darwin":
        return None
    if not shutil.which("osascript"):
        return "osascript not found"
    # Close the document even when the save fails, so Word isn't left holding
    # it open (with a ~$ lock file in the output folder). Word is sandboxed on
    # macOS: saving into a folder it has no access to (e.g. /tmp) fails with a
    # misleading "doesn't understand the save as message" (-1708).
    script = f'''
tell application "Microsoft Word"
    open POSIX file "{_applescript_str(docx_abs)}"
    set theDoc to active document
    try
        save as theDoc file name "{_applescript_str(pdf_abs)}" file format format PDF
    on error errMsg number errNum
        close theDoc saving no
        error errMsg number errNum
    end try
    close theDoc saving no
end tell
'''
    result = subprocess.run(['osascript', '-e', script], capture_output=True, text=True)
    if result.returncode == 0 and os.path.exists(pdf_abs):
        return None
    return result.stderr.strip() or f"osascript exited {result.returncode}"


def _export_word_windows(docx_abs, pdf_abs):
    """Word via docx2pdf (COM automation)."""
    if sys.platform != "win32":
        return None
    try:
        from docx2pdf import convert
    except ImportError:
        return "docx2pdf not installed (pip install docx2pdf)"
    try:
        convert(docx_abs, pdf_abs)
    except Exception as e:
        return str(e)
    return None


def _find_soffice():
    import shutil

    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    candidates = []
    if sys.platform == "darwin":
        candidates.append("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    elif sys.platform == "win32":
        for env in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
            base = os.environ.get(env)
            if base:
                candidates.append(os.path.join(base, "LibreOffice", "program", "soffice.exe"))
    for c in candidates:
        if os.path.exists(c):
            return c
    return None


def _export_libreoffice(docx_abs, pdf_abs):
    """LibreOffice headless. Writes <basename>.pdf into the output folder."""
    import subprocess

    soffice = _find_soffice()
    if not soffice:
        return "LibreOffice (soffice) not found"
    out_dir = os.path.dirname(pdf_abs)
    try:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_abs],
            capture_output=True, text=True, timeout=180
        )
    except (OSError, subprocess.TimeoutExpired) as e:
        return str(e)
    stem = os.path.splitext(os.path.basename(docx_abs))[0]
    lo_pdf = os.path.join(out_dir, stem + ".pdf")
    if os.path.exists(lo_pdf) and lo_pdf != pdf_abs:
        os.replace(lo_pdf, pdf_abs)
    if os.path.exists(pdf_abs):
        return None
    return result.stderr.strip() or f"soffice exited {result.returncode}"


def _export_pdf(docx_path, pdf_path):
    """
    Convert .docx to .pdf. Tries backends in order and returns the name of the
    one that produced the PDF, or None if all failed:
      1. Word — osascript on macOS, docx2pdf on Windows.
         Word renders justified text correctly; LibreOffice does not.
         macOS will prompt for Automation access the very first time only.
      2. LibreOffice headless — fallback on every platform.
    Any stale PDF at pdf_path is removed first so an old file never counts.
    """
    docx_abs = os.path.abspath(docx_path)
    pdf_abs  = os.path.abspath(pdf_path)
    if os.path.exists(pdf_abs):
        try:
            os.remove(pdf_abs)
        except OSError as e:
            print(f"⚠ PDF export failed: can't replace the existing PDF ({e.strerror}).")
            print(f"  {pdf_abs}")
            print("  → Close it in any PDF viewer or Word, then re-run.")
            return None

    errors = []
    word = _export_word_mac if sys.platform == "darwin" else _export_word_windows
    for name, backend in (("Word", word), ("LibreOffice", _export_libreoffice)):
        err = backend(docx_abs, pdf_abs)
        if err is None and os.path.exists(pdf_abs):
            if errors:
                # Word was tried first and failed; say so, since LibreOffice
                # renders justified text differently.
                print(f"⚠ Word export failed, used LibreOffice instead ({errors[0]})")
            print(f"✓ PDF ({name}): {pdf_abs}")
            return name
        if err:
            errors.append(f"{name}: {err}")

    print("⚠ PDF export failed (no Word or LibreOffice backend produced a PDF).")
    for e in errors:
        print(f"  {e}")
    print("  → Install Microsoft Word, or LibreOffice (libreoffice.org), and re-run.")
    return None


def count_pdf_pages(pdf_path):
    """Page count via pypdf, or None if pypdf is not installed."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return None
    return len(PdfReader(pdf_path).pages)


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Generate a formatted .docx (and .pdf) resume from a JSON content file.")
    parser.add_argument("content", help="resume content JSON file")
    parser.add_argument("output", help='output filename, e.g. "Your Name Resume MM.DD.YYYY Company Role.docx"')
    parser.add_argument("--pages", type=int, default=2,
                        help="required PDF page count (default: 2); exit 1 on mismatch")
    parser.add_argument("--no-pdf", action="store_true",
                        help="write the .docx only; skip PDF export and the page check")
    args = parser.parse_args()

    content_path = args.content
    output_name  = args.output

    with open(content_path, "r", encoding="utf-8") as f:
        content = json.load(f)

    master_path = _find_master()
    doc = Document(master_path)
    preflight_master(doc, master_path)

    # Dates tab stop sits on the master's right margin
    global RIGHT_TAB_POS
    RIGHT_TAB_POS = right_tab_pos(doc)

    # Kill auto-hyphenation inherited from MASTER settings
    disable_auto_hyphenation(doc)

    # Update only the headline text — leave name, contact, links untouched
    update_headline_in_place(doc, content.get("headline", ""))

    # Clear body paragraphs (preserves header table and sectPr)
    clear_document_body(doc)

    build_summary(doc, content.get("summary", ""))
    build_experience(doc, content.get("experience", []))
    build_skills(doc, content.get("skills", []))
    build_education(doc, content.get("education", ""))

    if not output_name.lower().endswith(".docx"):
        output_name += ".docx"
    output_path = os.path.join(OUTPUT_DIR, output_name)

    # Swap Word-bundled Garamond for EB Garamond (installed system font) so that
    # both the DOCX and the LibreOffice-generated PDF use the same font.
    # EB Garamond is a high-quality Garamond-family font accessible to both Word
    # and LibreOffice. Install from: fonts.google.com/specimen/EB+Garamond
    _replace_font(doc, "Garamond", "EB Garamond")

    doc.save(output_path)
    print(f"✓ Saved: {output_path}")

    if args.no_pdf:
        print("⚠ PDF export and page check skipped (--no-pdf).")
        return

    # Auto-export PDF alongside the .docx
    pdf_path = os.path.splitext(output_path)[0] + ".pdf"
    if _export_pdf(output_path, pdf_path) is None:
        sys.exit(1)

    # Verify page count
    pages = count_pdf_pages(pdf_path)
    if pages is None:
        print("⚠ Page check failed: pypdf is not installed (pip install pypdf).")
        sys.exit(1)
    if pages > args.pages:
        print(f"⚠ Page count: {pages} (expected {args.pages}): over by {pages - args.pages}. "
              "Cut or tighten content and regenerate.")
        sys.exit(1)
    if pages < args.pages:
        print(f"⚠ Page count: {pages} (expected {args.pages}): under by {args.pages - pages}. "
              "Add content and regenerate.")
        sys.exit(1)
    print(f"✓ Page count: {pages}")


if __name__ == "__main__":
    main()
