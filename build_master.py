#!/usr/bin/env python3
"""
build_master.py — Write a starter master_resume.docx for generate_resume.py.

Use this if you don't have a master resume .docx of your own. The file it
writes has exactly the structure the generator expects:

  - A one-cell header table:
      paragraph 0: Name
      paragraph 1: Headline + <w:br/> + Contact
      paragraph 2: Links
  - Paragraph styles "Body A", "Section Title" and "List Paragraph"
  - A bullet list definition attached to "List Paragraph"
  - Letter paper, 0.8" margins, Garamond 11pt

Usage:
    python3 build_master.py --name "Jane Smith" \\
        --contact "555-555-5555 • jane@example.com • City, ST" \\
        --links "linkedin.com/in/janesmith • janesmith.com"

Every argument is optional; placeholders are used for anything you leave out.
Open the result in Word to adjust fonts, sizes or spacing — the generator
inherits whatever you change. Refuses to overwrite an existing file unless
--force is given.
"""

import sys
import os
import argparse
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Constants ────────────────────────────────────────────────────────────────

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(SCRIPT_DIR, "master_resume.docx")

FONT = "Garamond"   # generate_resume.py swaps this for EB Garamond at build time
BODY_PT = 11

PLACEHOLDER_NAME = "Your Full Name"
PLACEHOLDER_HEADLINE = "Your Headline | Replaced On Every Resume"
PLACEHOLDER_CONTACT = "(555) 555-5555 • you@example.com • City, ST"
PLACEHOLDER_LINKS = "linkedin.com/in/yourname • yourportfolio.com"


# ── Style helpers ────────────────────────────────────────────────────────────

def set_font(style, size_pt, bold=None):
    """Set font name on all four rFonts slots so Word and LibreOffice agree."""
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), FONT)


def tight_spacing(style, before=0, after=0):
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0


def add_bottom_border(style):
    """Thin rule under section titles."""
    pPr = style.element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'auto')
    pBdr.append(bottom)
    pPr.append(pBdr)


def get_or_add_style(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


# ── Bullet numbering ─────────────────────────────────────────────────────────

def add_bullet_numbering(doc):
    """
    Append a single-level bullet definition to numbering.xml and return its numId.
    The generator reads the numId back from the List Paragraph style.
    """
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(a.get(qn('w:abstractNumId')))
                    for a in numbering.findall(qn('w:abstractNum'))]
    num_ids = [int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]
    abstract_id = str(max(abstract_ids, default=-1) + 1)
    num_id = str(max(num_ids, default=0) + 1)

    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), abstract_id)
    multi = OxmlElement('w:multiLevelType')
    multi.set(qn('w:val'), 'singleLevel')
    abstract.append(multi)

    lvl = OxmlElement('w:lvl')
    lvl.set(qn('w:ilvl'), '0')
    for tag, val in (('w:start', '1'), ('w:numFmt', 'bullet'),
                     ('w:lvlText', '•'), ('w:lvlJc', 'left')):
        el = OxmlElement(tag)
        el.set(qn('w:val'), val)
        lvl.append(el)
    pPr = OxmlElement('w:pPr')
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), '360')
    ind.set(qn('w:hanging'), '216')
    pPr.append(ind)
    lvl.append(pPr)
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rFonts.set(qn(attr), FONT)
    rPr.append(rFonts)
    lvl.append(rPr)
    abstract.append(lvl)

    num = OxmlElement('w:num')
    num.set(qn('w:numId'), num_id)
    ref = OxmlElement('w:abstractNumId')
    ref.set(qn('w:val'), abstract_id)
    num.append(ref)

    # Schema order: every abstractNum precedes every num
    first_num = numbering.find(qn('w:num'))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(num)
    return num_id


def attach_numbering(style, num_id):
    pPr = style.element.get_or_add_pPr()
    numPr = OxmlElement('w:numPr')
    numId = OxmlElement('w:numId')
    numId.set(qn('w:val'), num_id)
    numPr.append(numId)
    pPr.insert(0, numPr)


# ── Builders ─────────────────────────────────────────────────────────────────

def build_styles(doc):
    normal = doc.styles["Normal"]
    set_font(normal, BODY_PT)
    tight_spacing(normal)

    body = get_or_add_style(doc, "Body A")
    body.base_style = normal
    set_font(body, BODY_PT)
    tight_spacing(body)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = get_or_add_style(doc, "Section Title")
    title.base_style = normal
    set_font(title, 12, bold=True)
    title.font.all_caps = True
    tight_spacing(title, before=8, after=3)
    title.paragraph_format.keep_with_next = True
    add_bottom_border(title)

    bullet = doc.styles["List Paragraph"]
    bullet.base_style = normal
    set_font(bullet, BODY_PT)
    tight_spacing(bullet)
    bullet.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    bullet.paragraph_format.left_indent = None   # indent comes from the bullet level
    attach_numbering(bullet, add_bullet_numbering(doc))


def build_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(section, side, Inches(0.8))


def build_header(doc, name, headline, contact, links):
    """One-cell table: name / headline + <w:br/> + contact / links."""
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]

    p_name = cell.paragraphs[0]
    p_name.style = doc.styles["Body A"]
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_name.add_run(name)
    run.bold = True
    run.font.size = Pt(18)

    p_head = cell.add_paragraph(style="Body A")
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_head.add_run(headline).bold = True
    br_run = p_head.add_run()
    br_run.add_break()
    br_run.add_text(contact)

    p_links = cell.add_paragraph(style="Body A")
    p_links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_links.add_run(links)


def build_placeholder_body(doc):
    """Sample body so the file looks like a resume. The generator clears it."""
    doc.add_paragraph("Summary", style="Section Title")
    doc.add_paragraph(
        "This body text is a placeholder. generate_resume.py replaces everything "
        "below the header table with content from your JSON file.",
        style="Body A")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Write a starter master_resume.docx for generate_resume.py.")
    parser.add_argument("--name", default=PLACEHOLDER_NAME)
    parser.add_argument("--contact", default=PLACEHOLDER_CONTACT,
                        help="phone • email • city line shown under the headline")
    parser.add_argument("--links", default=PLACEHOLDER_LINKS,
                        help="LinkedIn / portfolio line")
    parser.add_argument("--force", action="store_true",
                        help="overwrite an existing master_resume.docx")
    args = parser.parse_args()

    if os.path.exists(OUTPUT_PATH) and not args.force:
        print(f"⚠ {OUTPUT_PATH} already exists. Re-run with --force to overwrite it.")
        sys.exit(1)

    doc = Document()
    build_page(doc)
    build_styles(doc)
    build_header(doc, args.name, PLACEHOLDER_HEADLINE, args.contact, args.links)
    build_placeholder_body(doc)

    doc.save(OUTPUT_PATH)
    print(f"✓ Saved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
